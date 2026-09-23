from __future__ import annotations
from pathlib import Path
import re
import zipfile
import pandas as pd
from docx import Document

IMAGE_EXTS = {'.png','.jpg','.jpeg','.webp','.bmp','.tif','.tiff'}
EXCEL_EXTS = {'.xlsx','.xlsm','.xls'}


def normalize_sheet(name: str) -> str:
    return re.sub(r'\s+', '_', str(name).strip().lower())


def validate_excel(path: str | Path, required_any=None, label='Excel workbook') -> dict:
    p = Path(path)
    if p.suffix.lower() not in EXCEL_EXTS:
        raise ValueError(f'{label} must be an Excel workbook (.xlsx/.xlsm/.xls): {p.name}')
    try:
        xls = pd.ExcelFile(p)
    except Exception as e:
        raise ValueError(f'{label} cannot be opened as Excel: {type(e).__name__}: {e}') from e
    sheets = list(xls.sheet_names)
    if required_any:
        normalized = {normalize_sheet(s) for s in sheets}
        wanted = {normalize_sheet(s) for s in required_any}
        if not normalized.intersection(wanted):
            raise ValueError(f'{label} does not contain any expected sheet. Found: {", ".join(sheets)}')
    return {'valid': True, 'type': 'excel', 'name': p.name, 'sheets': sheets}


def validate_raw_ar(path: str | Path) -> dict:
    return validate_excel(path, {
        'NEW_ITEMS','NEW ITEMS','NOT_INCLUDED','NOT INCLUDED',
        'DROPPED_ITEMS','DROPPED ITEMS','DROPPED_UPCS','DROPPED UPCS',
        'CHANGED_ITEMS','CHANGED ITEMS'
    }, 'Raw AR workbook')


def validate_docx(path: str | Path, label='DB Guide') -> dict:
    p = Path(path)
    if p.suffix.lower() != '.docx':
        raise ValueError(f'{label} must be a DOCX file: {p.name}')
    try:
        doc = Document(p)
        text = '\n'.join(x.text for x in doc.paragraphs if x.text.strip())
        table_text = '\n'.join(' | '.join(str(c.text or '') for c in row.cells) for t in doc.tables for row in t.rows)
        total = (text + '\n' + table_text).strip()
    except Exception as e:
        raise ValueError(f'{label} cannot be opened as DOCX: {type(e).__name__}: {e}') from e
    if len(total) < 50:
        raise ValueError(f'{label} appears empty or unreadable.')
    return {'valid': True, 'type': 'docx', 'name': p.name, 'characters': len(total)}


def validate_image(path: str | Path) -> dict:
    p = Path(path)
    if p.suffix.lower() not in IMAGE_EXTS:
        raise ValueError(f'Unsupported image format: {p.name}')
    try:
        from PIL import Image
        with Image.open(p) as im:
            im.verify()
        with Image.open(p) as im:
            size = im.size
    except Exception as e:
        raise ValueError(f'Image is invalid or unreadable: {p.name}: {type(e).__name__}: {e}') from e
    return {'valid': True, 'type': 'image', 'name': p.name, 'size': list(size)}


def validate_image_zip(path: str | Path, max_files=100000, max_uncompressed_mb=4096) -> dict:
    p = Path(path)
    if p.suffix.lower() != '.zip':
        raise ValueError(f'Image package must be a ZIP file: {p.name}')
    try:
        with zipfile.ZipFile(p) as z:
            bad = z.testzip()
            if bad:
                raise ValueError(f'Image ZIP is corrupt near: {bad}')
            infos = [i for i in z.infolist() if not i.is_dir()]
            images = [i for i in infos if Path(i.filename).suffix.lower() in IMAGE_EXTS]
            if not images:
                raise ValueError('Image ZIP contains no supported image files.')
            if len(infos) > max_files:
                raise ValueError(f'Image ZIP contains too many files ({len(infos)}). Limit: {max_files}.')
            total = sum(max(0, int(i.file_size)) for i in infos)
            if total > max_uncompressed_mb * 1024 * 1024:
                raise ValueError(f'Image ZIP uncompressed size exceeds {max_uncompressed_mb} MB.')
    except zipfile.BadZipFile as e:
        raise ValueError(f'Invalid ZIP file: {p.name}') from e
    return {'valid': True, 'type': 'image_zip', 'name': p.name, 'image_files': len(images), 'total_files': len(infos), 'uncompressed_bytes': total}
