from pathlib import Path
import openpyxl
import re

BASE = Path(__file__).resolve().parent

def norm(v):
    if v is None: return ''
    s = str(v).strip()
    return re.sub(r'\.0$', '', s)

def key(ws, row, bucket):
    headers = {str(ws.cell(1,c).value).strip(): c for c in range(1, ws.max_column+1) if ws.cell(1,c).value is not None}
    if bucket == 'NOT_INCLUDED':
        nc = headers.get('NAN_KEY', headers.get('NANKEY'))
        uc = headers.get('BARCODE', headers.get('UPC'))
        dc = headers.get('PRO ITEM DESCRIPTION ', headers.get('PRO ITEM DESCRIPTION', headers.get('ITEM_DESCRIPTION')))
        return ('NOT_INCLUDED', norm(ws.cell(row, nc).value if nc else '').upper(),
                norm(ws.cell(row, uc).value if uc else ''),
                norm(ws.cell(row, dc).value if dc else '').upper())
    nc = headers.get('NANKEY', headers.get('NAN_KEY', headers.get('NAN KEY')))
    uc = headers.get('UPC', headers.get('BARCODE'))
    dc = headers.get('PRO ITEM DESCRIPTION ', headers.get('PRO ITEM DESCRIPTION', headers.get('ITEM DESCRIPTION')))
    nk = norm(ws.cell(row, nc).value if nc else '')
    upc = norm(ws.cell(row, uc).value if uc else '')
    desc = norm(ws.cell(row, dc).value if dc else '').upper()
    return (bucket,nk,upc,desc)

def check(output, reference):
    owb=openpyxl.load_workbook(output,data_only=True)
    rwb=openpyxl.load_workbook(reference,data_only=True)
    checks=[]
    for bucket, sheet in [('NEW_ITEMS','NEW_ITEMS'),('NOT_INCLUDED','NOT_INCLUDED'),('DROPPED_UPCS','DROPPED_UPCS')]:
        if sheet not in owb.sheetnames or sheet not in rwb.sheetnames: continue
        o=owb[sheet]; r=rwb[sheet]
        ref_index={}
        rh={str(r.cell(1,c).value).strip():c for c in range(1,r.max_column+1) if r.cell(1,c).value is not None}
        for rr in range(2,r.max_row+1):
            k=key(r,rr,bucket)
            ref_index[k]=rr
        oh={str(o.cell(1,c).value).strip():c for c in range(1,o.max_column+1) if o.cell(1,c).value is not None}
        for rr in range(2,o.max_row+1):
            if all(o.cell(rr,c).value is None for c in range(1,o.max_column+1)): continue
            k=key(o,rr,bucket)
            if k in ref_index:
                checks.append((bucket,rr,'AI COMMENTS',o.cell(rr,oh['AI COMMENTS']).value,r.cell(ref_index[k],rh['AI COMMENTS']).value))
    mism=[x for x in checks if x[3]!=x[4]]
    print(f"Reference comment checks: {len(checks)}")
    print(f"Comment mismatches: {len(mism)}")
    for x in mism[:20]: print(x)
    return 0 if not mism else 1

if __name__ == '__main__':
    import sys
    if len(sys.argv) != 3:
        print('Usage: python validate_output_against_reference.py OUTPUT.xlsx REFERENCE.xlsx')
        raise SystemExit(2)
    raise SystemExit(check(sys.argv[1],sys.argv[2]))
